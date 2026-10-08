"""Tạo tệp mẫu thật (Word, Excel, PowerPoint, PDF, ảnh) để kiểm thử bộ đọc tệp."""
import pathlib
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw

out = pathlib.Path(__file__).resolve().parent.parent / "tests/fixtures"; out.mkdir(exist_ok=True, parents=True)

d = Document(); d.add_heading("Báo cáo thử nghiệm", 1)
d.add_paragraph("Doanh thu quý 3 đạt 1.250 triệu đồng, tăng 12% so với quý 2.")
p = d.add_paragraph("Đoạn có "); p.add_run("chữ đậm").bold = True; p.add_run(" & ký tự đặc biệt <tag>.")
t = d.add_table(rows=2, cols=2); t.cell(0,0).text="Khu vực"; t.cell(0,1).text="Doanh thu"; t.cell(1,0).text="Miền Nam"; t.cell(1,1).text="700"
d.save(out/"mau.docx")

wb = Workbook(); ws = wb.active; ws.title = "Doanh thu"
ws.append(["Tháng","Doanh thu","Ghi chú"]); ws.append(["Một",100,"ổn, tốt"]); ws.append(["Hai",150.5,'có "trích dẫn"']); ws.append(["Ba",True,None])
ws2 = wb.create_sheet("Chi phí"); ws2.append(["Khoản","Số tiền"]); ws2.append(["Thuê",30])
wb.save(out/"mau.xlsx")

pr = Presentation()
for title, body in [("Giới thiệu","Mục tiêu năm nay"),("Kết quả","Tăng trưởng 12%")]:
    s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = title; s.placeholders[1].text = body
pr.save(out/"mau.pptx")

c = canvas.Canvas(str(out/"mau.pdf"), pagesize=A4); c.drawString(72, 750, "Hop dong thu nghiem so 42"); c.showPage(); c.save()

img = Image.new("RGB", (2400, 1600), "#f4efe6"); dr = ImageDraw.Draw(img)
dr.rectangle([150,150,2250,1450], outline="#0b7c86", width=40)
dr.ellipse([400,350,1100,1050], fill="#e3563b"); dr.rectangle([1250,350,2000,1050], fill="#0b7c86"); dr.polygon([(700,1400),(1200,1100),(1700,1400)], fill="#f2b705")
img.save(out/"anh-lon.png"); img.resize((300,200)).save(out/"anh-nho.jpg", quality=90)
(out/"ghi-chu.txt").write_text("Nội dung tệp văn bản: số bí mật là 42.", encoding="utf-8")
print("Đã tạo:", sorted(p.name for p in out.iterdir()))
